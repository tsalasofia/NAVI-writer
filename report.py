"""Report writer: templates -> section-by-section RAG generation -> .docx"""
import datetime
import io
import json
import re
from pathlib import Path

import ollama
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from rag import CHAT_MODEL, NUM_CTX, doc_collection, embed_query

TEMPLATE_DIR = Path("templates")
REPORT_DIR = Path("reports")

SECTION_TOP_K = 6              # chunks retrieved per section
SECTION_CONTEXT_CHARS = 6000   # max source text per section
REPORT_SO_FAR_CHARS = 5000     # max text of the rest of the report for summary sections
NO_MATERIAL_DISTANCE = 0.9     # if even the best chunk is farther than this -> no material

# ======================================================================
# Templates
# ======================================================================
# A template is a JSON file in ./templates describing structure and style.
# Section keys:
#   title, guidance (what to write), target_words,
#   style: "prose" | "bullets" | "mixed",
#   from_report: true -> written LAST, from the other sections
#                (executive summary, conclusion, abstract ...)

DEFAULT_TEMPLATES = {
    "school_report": {
        "name": "School Report",
        "description": "Formal academic report with title page, numbered references and 1.5 line spacing.",
        "tone": "formal, clear, objective, academic, third person",
        "font": "Times New Roman",
        "font_size": 12,
        "line_spacing": 1.5,
        "heading_color": "000000",
        "title_page": True,
        "references": True,
        "builtin": True,
        "sections": [
            {"title": "Introduction", "target_words": 150, "style": "prose",
             "guidance": "Introduce the topic, why it matters, and what the report covers."},
            {"title": "Background", "target_words": 250, "style": "prose",
             "guidance": "Give the context, key concepts and history needed to understand the topic."},
            {"title": "Main Discussion", "target_words": 450, "style": "prose",
             "guidance": "Present the main facts, arguments and evidence from the sources in a logical order."},
            {"title": "Conclusion", "target_words": 150, "style": "prose", "from_report": True,
             "guidance": "Summarise the main points and give a closing statement. Introduce no new facts."},
        ],
    },
    "business_report": {
        "name": "Business Report",
        "description": "Concise professional report with executive summary, findings and recommendations.",
        "tone": "professional, concise, direct, results-oriented",
        "font": "Calibri",
        "font_size": 11,
        "line_spacing": 1.15,
        "heading_color": "1F3864",
        "title_page": False,
        "references": True,
        "builtin": True,
        "sections": [
            {"title": "Executive Summary", "target_words": 150, "style": "prose", "from_report": True,
             "guidance": "A short summary of the purpose, key findings and main recommendations for busy readers."},
            {"title": "Purpose and Scope", "target_words": 120, "style": "prose",
             "guidance": "State what this report is for and what it covers."},
            {"title": "Current Situation", "target_words": 250, "style": "prose",
             "guidance": "Describe the current state of affairs using facts from the sources."},
            {"title": "Key Findings", "target_words": 300, "style": "mixed",
             "guidance": "The most important facts and insights, as short bullets with brief explanations."},
            {"title": "Analysis", "target_words": 300, "style": "prose",
             "guidance": "Interpret the findings: causes, risks, opportunities and implications."},
            {"title": "Recommendations", "target_words": 220, "style": "bullets",
             "guidance": "Concrete, actionable recommendations that follow from the findings."},
            {"title": "Conclusion", "target_words": 100, "style": "prose", "from_report": True,
             "guidance": "Wrap up the report in a few sentences."},
        ],
    },
}


def slugify(text):
    return re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_") or "item"


def ensure_default_templates():
    TEMPLATE_DIR.mkdir(exist_ok=True)
    for tid, tpl in DEFAULT_TEMPLATES.items():
        path = TEMPLATE_DIR / f"{tid}.json"
        if not path.exists():
            path.write_text(json.dumps(tpl, indent=2), encoding="utf-8")


def list_templates():
    """Returns {template_id: template_dict}."""
    ensure_default_templates()
    out = {}
    for path in sorted(TEMPLATE_DIR.glob("*.json")):
        try:
            out[path.stem] = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue
    return out


def save_template(template):
    TEMPLATE_DIR.mkdir(exist_ok=True)
    tid = slugify(template["name"])
    (TEMPLATE_DIR / f"{tid}.json").write_text(
        json.dumps(template, indent=2), encoding="utf-8"
    )
    return tid


def delete_template(template_id):
    path = TEMPLATE_DIR / f"{template_id}.json"
    if path.exists():
        path.unlink()


def _base_template(name, description, tone, font, font_size, title_page):
    return {
        "name": name,
        "description": description,
        "tone": tone,
        "font": font,
        "font_size": font_size,
        "line_spacing": 1.15,
        "heading_color": "000000",
        "title_page": title_page,
        "references": True,
        "builtin": False,
        "sections": [],
    }


def template_from_outline(name, outline, description="", tone="professional and clear",
                          font="Calibri", font_size=11, title_page=False):
    """Outline format, one section per line:  Title | guidance | words(optional)
    Add '*' at the end of the title to mark a section written from the rest of
    the report (summary / conclusion)."""
    tpl = _base_template(name, description, tone, font, font_size, title_page)
    for line in outline.splitlines():
        line = line.strip()
        if not line:
            continue
        parts = [p.strip() for p in line.split("|")]
        title = parts[0]
        from_report = title.endswith("*")
        title = title.rstrip("*").strip()
        guidance = parts[1] if len(parts) > 1 else ""
        words = 250
        if len(parts) > 2 and parts[2].isdigit():
            words = int(parts[2])
        tpl["sections"].append({
            "title": title, "guidance": guidance, "target_words": words,
            "style": "prose", "from_report": from_report,
        })
    return tpl


def template_from_docx(name, data: bytes, **kwargs):
    """Reads a .docx: every Heading becomes a section and the text written
    under it becomes the guidance for that section. Also copies the body font."""
    doc = Document(io.BytesIO(data))
    tpl = _base_template(
        name, kwargs.get("description", "Imported from a Word document."),
        kwargs.get("tone", "professional and clear"),
        doc.styles["Normal"].font.name or "Calibri",
        int(doc.styles["Normal"].font.size.pt) if doc.styles["Normal"].font.size else 11,
        False,
    )
    current = None
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        if p.style.name.lower().startswith("heading"):
            current = {"title": text, "guidance": "", "target_words": 250,
                       "style": "prose", "from_report": False}
            tpl["sections"].append(current)
        elif current is not None:
            current["guidance"] = (current["guidance"] + " " + text).strip()
    return tpl


# ======================================================================
# Retrieval for a section
# ======================================================================
class RefRegistry:
    """Maps (file, page) -> citation number."""

    def __init__(self):
        self.by_key = {}
        self.by_num = {}

    def number_for(self, file_name, page):
        key = (file_name, page)
        if key not in self.by_key:
            n = len(self.by_key) + 1
            self.by_key[key] = n
            self.by_num[n] = key
        return self.by_key[key]


def retrieve_for_section(query, file_names, registry):
    if doc_collection.count() == 0 or not file_names:
        return ""
    res = doc_collection.query(
        query_embeddings=[embed_query(query)],
        n_results=min(SECTION_TOP_K, doc_collection.count()),
        where={"file_name": {"$in": list(file_names)}},
        include=["documents", "metadatas", "distances"],
    )
    docs = res["documents"][0]
    metas = res["metadatas"][0]
    dists = res["distances"][0]
    if not docs or min(dists) > NO_MATERIAL_DISTANCE:
        return ""

    parts, used = [], 0
    for doc, meta in zip(docs, metas):
        n = registry.number_for(meta["file_name"], meta.get("page"))
        block = f"[{n}] ({meta['file_name']}" + (
            f", p.{meta['page']})\n" if meta.get("page") else ")\n"
        ) + doc
        if used + len(block) > SECTION_CONTEXT_CHARS:
            break
        parts.append(block)
        used += len(block)
    return "\n\n".join(parts)


# ======================================================================
# Writing sections with the LLM
# ======================================================================
def _clean_model_text(text, section_title):
    text = text.strip()
    lines = text.splitlines()
    # drop chatty preambles and repeated headings
    while lines and (
        re.match(r"^(here is|here's|sure|certainly|below is)", lines[0].strip(), re.I)
        or lines[0].strip().lstrip("#* ").rstrip(":*").lower() == section_title.lower()
        or not lines[0].strip()
    ):
        lines.pop(0)
    return "\n".join(lines).strip()


def write_section(section, topic, tone, context, report_so_far, words):
    title = section["title"]
    style = section.get("style", "prose")
    style_rule = {
        "prose": "Write in flowing paragraphs. Do not use bullet points.",
        "bullets": "Write as a bulleted list; every line starts with '- '.",
        "mixed": "Use a short intro sentence followed by '- ' bullet points.",
    }.get(style, "Write in flowing paragraphs.")

    system = (
        "You are a professional report writer.\n"
        f"Tone: {tone}. Write in English.\n"
        f"Write ONLY the body text of the section '{title}'. No heading, no preamble, no closing remarks.\n"
        "Use ONLY facts found in the material provided. Never invent facts, numbers, names or quotes.\n"
        "When you use a fact from a numbered source excerpt, cite it as [n] right after the sentence.\n"
        f"{style_rule}\n"
        f"Target length: about {words} words."
    )

    material = ""
    if report_so_far:
        material += "REPORT SO FAR (summarise or draw on it):\n" + report_so_far + "\n\n"
    if context:
        material += "SOURCE EXCERPTS:\n" + context
    user = (
        f"Report topic / instructions: {topic}\n\n"
        f"Section to write: {title}\n"
        f"What this section should contain: {section.get('guidance') or 'Use your judgment.'}\n\n"
        f"{material}\n\nWrite the section now."
    )

    res = ollama.chat(
        model=CHAT_MODEL,
        messages=[{"role": "system", "content": system},
                  {"role": "user", "content": user}],
        options={"num_ctx": NUM_CTX, "temperature": 0.3},
    )
    return _clean_model_text(res["message"]["content"], title)


# ======================================================================
# Citation renumbering
# ======================================================================
CITE_RE = re.compile(r"\[(\d+(?:\s*,\s*\d+)*)\]")


def renumber_citations(section_texts, registry):
    """Renumbers citations by order of first appearance, drops invalid ones,
    and returns (new_texts, reference_lines)."""
    order = {}

    def collect(m):
        for n in (int(x) for x in m.group(1).split(",")):
            if n in registry.by_num and n not in order:
                order[n] = len(order) + 1
        return m.group(0)

    for t in section_texts:
        CITE_RE.sub(collect, t)

    def replace(m):
        nums = [int(x) for x in m.group(1).split(",")]
        new = sorted({order[n] for n in nums if n in order})
        return "[" + ", ".join(map(str, new)) + "]" if new else ""

    new_texts = [CITE_RE.sub(replace, t) for t in section_texts]

    refs = []
    for old, new in sorted(order.items(), key=lambda kv: kv[1]):
        file_name, page = registry.by_num[old]
        refs.append(f"[{new}] {file_name}" + (f", p. {page}" if page else ""))
    return new_texts, refs


# ======================================================================
# Building the .docx
# ======================================================================
def _set_font(style, font_name, size=None, color=None, bold=None):
    style.font.name = font_name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), font_name)
    if size:
        style.font.size = Pt(size)
    if color:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold


def _add_page_number(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _add_rich_text(paragraph, text):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if len(part) > 4 and part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part.replace("**", ""))


def _write_blocks(doc, text):
    """Converts light markdown (paragraphs, '- ' bullets, '1.' lists, '#' headings)."""
    buffer = []

    def flush():
        if buffer:
            _add_rich_text(doc.add_paragraph(), " ".join(buffer))
            buffer.clear()

    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            flush()
            continue
        m_head = re.match(r"^#{1,6}\s+(.*)", line)
        m_bul = re.match(r"^[-*•]\s+(.*)", line)
        m_num = re.match(r"^(\d+[.)])\s+(.*)", line)
        if m_head:
            flush()
            doc.add_heading(m_head.group(1).strip("* "), level=2)
        elif m_bul:
            flush()
            _add_rich_text(doc.add_paragraph(style="List Bullet"), m_bul.group(1))
        elif m_num:
            flush()
            # manual numbers: Word's "List Number" style keeps counting across sections
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            _add_rich_text(p, f"{m_num.group(1)} {m_num.group(2)}")
        else:
            buffer.append(line)
    flush()


def build_docx(template, title, author, sections, references):
    """sections: list of (title, text). Returns .docx bytes."""
    doc = Document()
    font = template.get("font", "Calibri")
    size = template.get("font_size", 11)
    color = template.get("heading_color", "000000")

    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1)
        s.left_margin = s.right_margin = Inches(1)

    _set_font(doc.styles["Normal"], font, size)
    doc.styles["Normal"].paragraph_format.line_spacing = template.get("line_spacing", 1.15)
    doc.styles["Normal"].paragraph_format.space_after = Pt(8)
    _set_font(doc.styles["Title"], font, 26, color, True)
    _set_font(doc.styles["Heading 1"], font, size + 5, color, True)
    _set_font(doc.styles["Heading 2"], font, size + 2, color, True)
    _set_font(doc.styles["List Bullet"], font, size)

    today = datetime.date.today().strftime("%B %d, %Y")

    if template.get("title_page"):
        for _ in range(6):
            doc.add_paragraph()
        t = doc.add_paragraph(title, style="Title")
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        meta = [template["name"]] + ([f"By {author}"] if author else []) + [today]
        for line in meta:
            p = doc.add_paragraph(line)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    else:
        doc.add_paragraph(title, style="Title")
        sub = doc.add_paragraph(" · ".join([template["name"]] + ([author] if author else []) + [today]))
        sub.runs[0].italic = True

    for sec_title, text in sections:
        doc.add_heading(sec_title, level=1)
        _write_blocks(doc, text)

    if template.get("references", True) and references:
        doc.add_heading("References", level=1)
        for ref in references:
            doc.add_paragraph(ref)

    _add_page_number(doc.sections[0])

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ======================================================================
# Orchestration
# ======================================================================
LENGTH_FACTORS = {"Short": 0.6, "Standard": 1.0, "Detailed": 1.6}


def generate_report(template, title, topic, file_names, author="",
                    length="Standard", progress=None):
    """Writes every section and returns
    {'docx': bytes, 'sections': [(title, text)], 'references': [str]}.
    progress(i, total, section_title) is called before each section."""
    factor = LENGTH_FACTORS.get(length, 1.0)
    tone = template.get("tone", "professional and clear")
    sections = template["sections"]
    registry = RefRegistry()
    texts = {}
    total = len(sections)

    # 1) normal sections first, 2) summary-type sections (from_report) last
    order = [i for i, s in enumerate(sections) if not s.get("from_report")]
    order += [i for i, s in enumerate(sections) if s.get("from_report")]

    for step, idx in enumerate(order, start=1):
        sec = sections[idx]
        if progress:
            progress(step, total, sec["title"])
        words = max(40, int(sec.get("target_words", 250) * factor))

        if sec.get("from_report"):
            body = "\n\n".join(
                f"{sections[i]['title']}:\n{texts[i]}" for i in range(len(sections)) if i in texts
            )[:REPORT_SO_FAR_CHARS]
            texts[idx] = write_section(sec, topic, tone, "", body, words)
        else:
            query = f"{topic}. {sec['title']}. {sec.get('guidance', '')}"
            context = retrieve_for_section(query, file_names, registry)
            if not context:
                texts[idx] = ("(No relevant information about this section was found "
                              "in the selected source files.)")
            else:
                texts[idx] = write_section(sec, topic, tone, context, "", words)

    ordered_titles = [s["title"] for s in sections]
    ordered_texts = [texts[i] for i in range(len(sections))]
    ordered_texts, refs = renumber_citations(ordered_texts, registry)

    docx_bytes = build_docx(template, title, author,
                            list(zip(ordered_titles, ordered_texts)), refs)
    return {"docx": docx_bytes,
            "sections": list(zip(ordered_titles, ordered_texts)),
            "references": refs}


def save_report_file(title, docx_bytes):
    REPORT_DIR.mkdir(exist_ok=True)
    stamp = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    path = REPORT_DIR / f"{slugify(title)}_{stamp}.docx"
    path.write_bytes(docx_bytes)
    return str(path)
